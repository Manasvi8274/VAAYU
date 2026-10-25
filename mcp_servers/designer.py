"""
"designer": Vaayu's own MCP server for MAKING things - images, logos, posters, diagrams, presentations and
documents. Everything on this laptop, free, saved in designs/ (user's rule: nothing outside this folder) and
opened on the screen.

  images   Stable Diffusion Turbo (models/hf, ~2.6 GB) on the GTX 1650: the image is drawn in half precision
           (1.7 GB of VRAM, ~9 s) and decoded in full precision - all-half made black images on this card.
           Free VRAM is checked first; if Ollama holds a model it is asked to let go (Vaayu reloads it, ~5 s).
           The engine is dropped after 5 idle minutes so it doesn't sit on 3 GB of RAM.
  diagrams the local model writes Mermaid, headless Edge draws it (mermaid.js in mcp_servers/vendor/mermaid)
  slides,  the local model writes the content (mcp_servers/common.py), python-pptx / python-docx lay it out
  documents

The image engine's libraries live in mcp_servers/vendor/designer-sd (transformers 4.x for diffusers; Vaayu's
own transformers 5 stays untouched) - so they're first on the path, before anything imports them.

    .venv/Scripts/python.exe -m mcp_servers.designer        # stdio server (Vaayu starts it)
"""

import json
import os
import re
import subprocess
import sys
import tempfile
import textwrap
import threading
import time
from datetime import datetime
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
SD_LIBS = REPO / "mcp_servers" / "vendor" / "designer-sd"
sys.path.insert(0, str(SD_LIBS))
sys.path.insert(1, str(REPO))

from mcp.server.mcpserver import MCPServer  # noqa: E402

from mcp_servers.common import OLLAMA, ask_json, said, slug  # noqa: E402

OUT = REPO / "designs"
MERMAID_JS = REPO / "mcp_servers" / "vendor" / "mermaid" / "node_modules" / "mermaid" / "dist" / "mermaid.min.js"
EDGE_PROFILE = REPO / "mcp_servers" / "vendor" / "mermaid" / "edge-profile"
IMAGE_MODEL = "stabilityai/sd-turbo"
STEPS = 3            # calibration knob: SD-Turbo is made for 1-4; 3 looked best for the time
IDLE_UNLOAD = 300    # seconds before the image engine lets go of its memory
SIZES = {"square": (512, 512), "portrait": (512, 768), "landscape": (768, 512)}
FONTS = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"

server = MCPServer("designer", instructions="Make images, logos, posters, diagrams, presentations and documents.")
_engine_lock = threading.Lock()
_engine: dict = {}
_last_use = 0.0


# --- the image engine --------------------------------------------------------------------------------------

def _edge() -> str:
    for base in (os.environ.get("ProgramFiles(x86)", ""), os.environ.get("ProgramFiles", ""),
                 r"C:\Program Files (x86)", r"C:\Program Files"):
        exe = Path(base) / "Microsoft" / "Edge" / "Application" / "msedge.exe"
        if exe.exists():
            return str(exe)
    raise RuntimeError("Microsoft Edge isn't installed, and I draw diagrams with it.")


def _free_gpu() -> None:
    import torch
    free, _total = torch.cuda.mem_get_info()
    if free > 2.3 * 2**30:
        return
    import urllib.request
    try:  # Ollama holds a model on the card: ask it to let go (Vaayu's next question loads his back)
        with urllib.request.urlopen(f"{OLLAMA}/api/ps", timeout=5) as r:
            loaded = [m["name"] for m in json.load(r).get("models", [])]
        for name in loaded:
            req = urllib.request.Request(f"{OLLAMA}/api/generate", json.dumps({"model": name, "keep_alive": 0}).encode(),
                                         {"Content-Type": "application/json"})
            urllib.request.urlopen(req, timeout=30).read()
        time.sleep(1.0)
    except OSError:
        pass


def _load():
    import torch
    from diffusers import AutoencoderKL, AutoPipelineForText2Image
    if not torch.cuda.is_available():
        raise RuntimeError("No graphics card available for drawing images.")
    pipe = AutoPipelineForText2Image.from_pretrained(IMAGE_MODEL, torch_dtype=torch.float16, variant="fp16",
                                                     safety_checker=None)
    pipe.enable_model_cpu_offload()
    vae = AutoencoderKL.from_pretrained(IMAGE_MODEL, subfolder="vae", torch_dtype=torch.float32)
    _engine.update(pipe=pipe, vae=vae)
    threading.Thread(target=_unload_when_idle, daemon=True, name="designer-idle").start()


def _unload_when_idle() -> None:
    while _engine:
        time.sleep(30)
        with _engine_lock:
            if _engine and time.time() - _last_use > IDLE_UNLOAD:
                _engine.clear()
                import gc
                import torch
                gc.collect()
                torch.cuda.empty_cache()
                return


def _draw(prompt: str, size: str = "square"):
    """A picture from words - half-precision drawing, full-precision decoding (black images otherwise)."""
    global _last_use
    import torch
    from PIL import Image
    width, height = SIZES.get(size, SIZES["square"])
    with _engine_lock:
        _free_gpu()
        if not _engine:
            _load()
        _last_use = time.time()
        latents = _engine["pipe"](prompt=prompt, negative_prompt=None, num_inference_steps=STEPS, guidance_scale=0.0,
                                  width=width, height=height, output_type="latent").images
        if torch.isnan(latents).any():  # the GTX 1650's half-precision fault (found: every 512x768 portrait came
            # out black) - drawn square instead and fitted to the shape asked for
            latents = _engine["pipe"](prompt=prompt, num_inference_steps=STEPS, guidance_scale=0.0, width=512,
                                      height=512, output_type="latent").images
        vae = _engine["vae"].to("cuda")
        try:
            with torch.no_grad():
                img = vae.decode(latents.to("cuda", torch.float32) / vae.config.scaling_factor).sample
        finally:
            vae.to("cpu")
            torch.cuda.empty_cache()
        _last_use = time.time()
    pixels = ((img.clamp(-1, 1) + 1) * 127.5).round().byte()[0].permute(1, 2, 0).cpu().numpy()
    picture = Image.fromarray(pixels)
    if picture.size != (width, height):  # the square fallback, cropped to the shape and scaled
        from PIL import ImageOps
        picture = ImageOps.fit(picture, (width, height), Image.LANCZOS)
    return picture


# --- saving and showing ------------------------------------------------------------------------------------

def _path(kind: str, name: str, ext: str) -> Path:
    OUT.mkdir(exist_ok=True)
    return OUT / f"{kind}-{slug(name)}-{datetime.now():%Y%m%d-%H%M%S}.{ext}"


def _done(what: str, path: Path, open_it: bool) -> str:
    if open_it:
        os.startfile(path)  # noqa: S606 - the user's own design, opened for them to see
    return f"Made {what} - saved as {path.relative_to(REPO).as_posix()}" + (", and opened it." if open_it else ".")


def _font(size: int, bold: bool = False):
    from PIL import ImageFont
    for name in (("segoeuib.ttf", "arialbd.ttf") if bold else ("segoeui.ttf", "arial.ttf")):
        if (FONTS / name).exists():
            return ImageFont.truetype(str(FONTS / name), size)
    return ImageFont.load_default()


def _wrap(draw, text: str, font, width: int) -> list[str]:
    lines = []
    for para in (text or "").split("\n"):
        line = ""
        for word in para.split():
            trial = f"{line} {word}".strip()
            if draw.textlength(trial, font=font) <= width:
                line = trial
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return [x for x in lines if x is not None]


# --- the tools ---------------------------------------------------------------------------------------------

@server.tool()
@said
def design_image(description: str, size: str = "square", open_it: bool = True) -> str:
    """Draw a picture from a description (on this laptop). size: square, portrait or landscape."""
    img = _draw(f"{description}, highly detailed, beautiful lighting, professional", size)
    path = _path("image", description, "png")
    img.save(path)
    return _done("the picture", path, open_it)


@server.tool()
@said
def design_logo(description: str, open_it: bool = True) -> str:
    """Design a logo: a symbol drawn from the description, with the name set underneath."""
    from PIL import Image, ImageDraw
    spec = ask_json('You design logos. JSON: {"name": the brand name exactly as the user said it (never shortened '
                    'to initials), "symbol": a few words describing the icon - no text in it, "color": main color '
                    'word}', description)
    name = str(spec.get("name") or description).strip()[:40]
    symbol = _draw(f"{spec.get('symbol') or description} logo icon, minimalist flat vector emblem, "
                   f"{spec.get('color', '')}, centered, simple shapes, clean white background, no text", "square")
    canvas = Image.new("RGB", (1024, 1200), "white")
    canvas.paste(symbol.resize((860, 860), Image.LANCZOS), (82, 50))
    draw = ImageDraw.Draw(canvas)
    font = _font(110 if len(name) < 14 else 70, bold=True)
    w = draw.textlength(name, font=font)
    draw.text(((1024 - w) / 2, 960), name, font=font, fill=(30, 30, 40))
    path = _path("logo", name, "png")
    canvas.save(path)
    return _done(f"a logo for {name}", path, open_it)


@server.tool()
@said
def design_poster(description: str, open_it: bool = True) -> str:
    """Design a poster or flyer: a drawn background with the title and details set on it."""
    from PIL import Image, ImageDraw
    spec = ask_json('You design posters. JSON: {"title": short headline, "details": 1-3 short lines (date, place, '
                    'call to action) separated by newlines, "background": a few words describing the background '
                    'picture - no text in it}', description)
    title, details = str(spec.get("title") or description)[:80], str(spec.get("details") or "")[:300]
    # A square picture on top, a solid panel for the words below: readable whatever the picture is, and square
    # is the size the GTX 1650 draws reliably (portrait came out black).
    art = _draw(f"{spec.get('background') or description}, vibrant illustration, no text", "square")
    poster = Image.new("RGB", (1024, 1536), _THEME["dark"])
    poster.paste(art.resize((1024, 1024), Image.LANCZOS), (0, 0))
    draw = ImageDraw.Draw(poster)
    draw.rectangle([0, 1024, 1024, 1034], fill=_THEME["accent"])
    y = 1060
    big = _font(84 if len(title) < 30 else 64, bold=True)
    for line in _wrap(draw, title, big, 900)[:2]:
        draw.text((62, y), line, font=big, fill="white")
        y += big.size + 18
    for line in _wrap(draw, details, _font(42), 900)[:5]:
        draw.text((62, y + 14), line, font=_font(42), fill=(225, 225, 235))
        y += 54
    path = _path("poster", title, "png")
    poster.save(path)
    return _done(f"the poster '{title}'", path, open_it)


def _render_mermaid(code: str, png: Path) -> str:
    """Draw Mermaid code with headless Edge; '' if fine, else Mermaid's error."""
    from PIL import Image, ImageChops
    html = (f'<!doctype html><html><head><meta charset="utf-8"><style>body{{margin:0;background:#fff}}#d{{padding:28px;'
            f'display:inline-block}}</style><script src="{MERMAID_JS.as_uri()}"></script></head><body><div id="d">'
            f'<pre class="mermaid">{code.replace("<", "&lt;")}</pre></div><script>mermaid.initialize({{startOnLoad:false,'
            f"theme:'default',securityLevel:'strict'}});mermaid.run().catch(e=>{{document.body.setAttribute("
            f"'data-error',String(e.message||e))}})</script></body></html>")
    with tempfile.TemporaryDirectory(dir=OUT) as tmp:
        page = Path(tmp) / "diagram.html"
        page.write_text(html, encoding="utf-8")
        base = [_edge(), "--headless=new", "--disable-gpu", "--allow-file-access-from-files", "--hide-scrollbars",
                f"--user-data-dir={EDGE_PROFILE}", "--virtual-time-budget=8000", "--window-size=1600,1200"]
        flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        dom = subprocess.run(base + ["--dump-dom", page.as_uri()], capture_output=True, text=True, timeout=60,
                             encoding="utf-8", errors="replace", creationflags=flags).stdout
        if m := re.search(r'data-error="([^"]*)"', dom):
            return m.group(1)[:300] or "syntax error"
        subprocess.run(base + [f"--screenshot={png}", page.as_uri()], capture_output=True, timeout=60, creationflags=flags)
    img = Image.open(png).convert("RGB")
    box = ImageChops.difference(img, Image.new("RGB", img.size, "white")).getbbox()
    if box:  # trim the empty page around the drawing
        img.crop((max(box[0] - 30, 0), max(box[1] - 30, 0), min(box[2] + 30, img.width), min(box[3] + 30, img.height))).save(png)
    return ""


@server.tool()
@said
def design_diagram(description: str, open_it: bool = True) -> str:
    """Draw a diagram - flowchart, process, mind map, timeline, org chart, sequence of steps."""
    system = ('You draw diagrams in Mermaid. JSON: {"title": short title, "mermaid": the Mermaid code}. Use flowchart TD '
              'unless a mindmap, timeline, sequenceDiagram or pie fits better. Keep labels short and put every '
              'label in double quotes, e.g. A["Wake up"] --> B{"Raining?"}.')
    spec = ask_json(system, description)
    path = _path("diagram", spec.get("title") or description, "png")
    error = _render_mermaid(str(spec.get("mermaid") or ""), path)
    if error:  # one more go, told what was wrong
        spec = ask_json(system, f"{description}\n\nYour last Mermaid code failed with: {error}\nFix it.")
        error = _render_mermaid(str(spec.get("mermaid") or ""), path)
        if error:
            raise RuntimeError(f"I couldn't get that diagram to draw ({error[:80]}).")
    path.with_suffix(".mmd").write_text(str(spec.get("mermaid")), encoding="utf-8")
    return _done(f"the diagram '{spec.get('title') or description}'", path, open_it)


_THEME = {"dark": (0x1B, 0x26, 0x3B), "accent": (0xF2, 0x8C, 0x28), "text": (0x22, 0x22, 0x22), "light": (0xF5, 0xF6, 0xFA)}


@server.tool()
@said
def design_presentation(description: str, slides: int = 6, open_it: bool = True) -> str:
    """Make a PowerPoint presentation on a topic: written, laid out and styled, with a drawn title picture."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    slides = max(3, min(int(slides or 6), 12))
    spec = ask_json('You write presentations. JSON: {"title": str, "subtitle": str, "slides": [{"title": str, '
                    f'"bullets": [3-5 short points]}}]}} with exactly {slides - 1} content slides.', description)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    rgb = {k: RGBColor(*v) for k, v in _THEME.items()}

    def box(slide, x, y, w, h, text, size, color, bold=False):
        frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        frame.word_wrap = True
        frame.text = text
        for p in frame.paragraphs:
            p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, color, "Segoe UI"
        return frame

    title = str(spec.get("title") or description)[:90]
    first = deck.slides.add_slide(blank)
    first.background.fill.solid()
    first.background.fill.fore_color.rgb = rgb["dark"]
    try:  # a drawn picture for the title slide - the deck is still made without one
        pic = _path("slide-art", title, "png")
        _draw(f"{title}, elegant illustration, wide composition, no text", "landscape").save(pic)
        first.shapes.add_picture(str(pic), Inches(7.2), Inches(0), height=Inches(7.5))
    except Exception:
        pass
    bar = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(2.4), Inches(0.15), Inches(2.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb["accent"]
    bar.line.fill.background()
    box(first, 1.0, 2.2, 6.0, 1.8, title, 40, RGBColor(255, 255, 255), bold=True)
    box(first, 1.0, 4.0, 6.0, 1.2, str(spec.get("subtitle") or ""), 20, RGBColor(220, 220, 230))
    for item in (spec.get("slides") or [])[: slides - 1]:
        s = deck.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb["light"]
        side = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), Inches(7.5))
        side.fill.solid()
        side.fill.fore_color.rgb = rgb["accent"]
        side.line.fill.background()
        box(s, 0.9, 0.5, 11.5, 1.0, str(item.get("title", ""))[:90], 34, rgb["dark"], bold=True)
        frame = box(s, 0.9, 1.8, 11.5, 5.2, "", 22, rgb["text"])
        for i, point in enumerate((item.get("bullets") or [])[:6]):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = f"\u2022  {point}"
            p.font.size, p.font.color.rgb, p.font.name = Pt(22), rgb["text"], "Segoe UI"
            p.space_after = Pt(12)
    path = _path("presentation", title, "pptx")
    deck.save(path)
    return _done(f"the presentation '{title}' ({len(deck.slides)} slides)", path, open_it)


@server.tool()
@said
def design_document(description: str, open_it: bool = True) -> str:
    """Write and lay out a Word document - resume, cover letter, letter, report, essay, notice, invitation."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    spec = ask_json('You write well-structured documents. JSON: {"title": str, "subtitle": str, "sections": '
                    '[{"heading": str, "paragraphs": [str], "bullets": [str]}]}. Leave paragraphs or bullets empty '
                    'when not needed. Use placeholders like [Your phone] for details you were not given.', description)
    doc = Document()
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(11)
    title = str(spec.get("title") or description)[:90]
    heading = doc.add_heading(title, level=0)
    heading.runs[0].font.color.rgb = RGBColor(*_THEME["dark"])
    if spec.get("subtitle"):
        doc.add_paragraph(str(spec["subtitle"])).runs[0].italic = True
    for section in spec.get("sections") or []:
        if section.get("heading"):
            doc.add_heading(str(section["heading"]), level=1)
        for para in section.get("paragraphs") or []:
            doc.add_paragraph(str(para))
        for point in section.get("bullets") or []:
            doc.add_paragraph(str(point), style="List Bullet")
    path = _path("document", title, "docx")
    doc.save(path)
    return _done(f"the document '{title}'", path, open_it)


@server.tool()
@said
def list_designs(count: int = 10) -> str:
    """The most recent things Vaayu designed."""
    files = sorted((p for p in OUT.glob("*.*") if p.suffix in (".png", ".pptx", ".docx")), key=lambda p: -p.stat().st_mtime)
    if not files:
        return "I haven't designed anything yet."
    return "Latest designs: " + "; ".join(p.name for p in files[: max(1, min(int(count), 30))])


if __name__ == "__main__":
    server.run("stdio")
